from datetime import UTC, datetime
from io import BytesIO

import pymupdf
from docx import Document
from fastapi import HTTPException, UploadFile

from app.modules.resume import repository
from app.modules.resume.optimization_repository import (
    get_optimization,
    save_optimization,
    update_optimization_review,
)
from app.modules.resume.schemas import ResumeReviewRequest
from app.workflows.resume_optimization_graph import resume_optimization_graph

from fastapi import HTTPException

from app.modules.resume.schemas import ResumeReviewRequest
from app.modules.resume.optimization_repository import (
    get_optimization,
    update_optimization_review,


)



from fastapi import HTTPException

from app.modules.resume.optimization_repository import (
    get_optimization,
)

from app.modules.resume.version_repository import (
    save_resume_version,
    get_resume_version,
)



MAX_FILE_SIZE = 5 * 1024 * 1024


async def upload_resume(file: UploadFile):
    filename = file.filename or ""

    if not filename.lower().endswith((".pdf", ".docx")):
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX files are supported",
        )

    content = await file.read(MAX_FILE_SIZE + 1)

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File size must not exceed 5 MB",
        )

    try:
        if filename.lower().endswith(".pdf"):
            with pymupdf.open(stream=content, filetype="pdf") as document:
                text = "\n".join(page.get_text() for page in document)
        else:
            document = Document(BytesIO(content))
            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text)

            text = "\n".join(paragraphs)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Unable to read the uploaded document",
        ) from exc

    if not text.strip():
        raise HTTPException(
            status_code=422,
            detail="No readable text found in the resume",
        )

    resume_data = {
        "filename": filename,
        "content": text.strip(),
        "uploaded_at": datetime.now(UTC),
    }

    await repository.save_resume(resume_data)
    return resume_data


async def get_resume():
    resume = await repository.get_resume()

    if resume is None:
        raise HTTPException(
            status_code=404,
            detail="Resume not found",
        )

    return resume


async def generate_optimization(job_description: str):
    resume = await get_resume()
    resume_text = resume["content"]

    if not resume_text.strip():
        raise HTTPException(
            status_code=422,
            detail="Saved resume has no extracted text",
        )

    result = await resume_optimization_graph.ainvoke({
        "resume_text": resume_text,
        "job_description": job_description,
        "optimization_result": None,
    })
    optimization = result["optimization_result"]

    document = await save_optimization(
        job_description=job_description,
        suggestions=[
            suggestion.model_dump()
            for suggestion in optimization.suggestions
        ],
        missing_skills=optimization.missing_skills,
        summary=optimization.summary,
        resume_text=resume_text,
    )

    return {
        "optimization_id": document["_id"],
        "status": document["status"],
        "suggestions": document["suggestions"],
        "missing_skills": document["missing_skills"],
        "summary": document["summary"],
    }


async def review_resume_optimization(
    optimization_id: str,
    request: ResumeReviewRequest,
):
    optimization = await get_optimization(optimization_id)

    if optimization is None:
        raise HTTPException(
            status_code=404,
            detail="Resume optimization not found",
        )

    if optimization["status"] != "pending":
        raise HTTPException(
            status_code=409,
            detail="Optimization has already been reviewed",
        )

    expected_ids = {
        suggestion["suggestion_id"]
        for suggestion in optimization["suggestions"]
    }
    submitted_ids = [item.suggestion_id for item in request.decisions]

    if len(submitted_ids) != len(set(submitted_ids)):
        raise HTTPException(
            status_code=422,
            detail="Duplicate suggestion IDs are not allowed",
        )

    decisions = {
        item.suggestion_id: item.decision
        for item in request.decisions
    }
    if set(decisions) != expected_ids:
        raise HTTPException(
            status_code=422,
            detail="Provide exactly one decision for every suggestion",
        )

    try:
        updated = await update_optimization_review(
            optimization_id=optimization_id,
            decisions=decisions,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=409,
            detail=str(exc),
        ) from exc

    if updated is None:
        raise HTTPException(
            status_code=404,
            detail="Resume optimization not found",
        )

    return {
        "optimization_id": updated["_id"],
        "status": updated["status"],
        "suggestions": updated["suggestions"],
        "missing_skills": updated["missing_skills"],
        "summary": updated["summary"],
    }





async def review_resume_optimization(
    optimization_id: str,
    request: ResumeReviewRequest,
):
    optimization = await get_optimization(optimization_id)

    if optimization is None:
        raise HTTPException(
            status_code=404,
            detail="Optimization not found",
        )

    if optimization["status"] != "pending":
        raise HTTPException(
            status_code=409,
            detail="Optimization already reviewed",
        )

    expected_ids = {
        suggestion["suggestion_id"]
        for suggestion in optimization["suggestions"]
    }

    submitted_ids = [
        item.suggestion_id
        for item in request.decisions
    ]

    if len(submitted_ids) != len(set(submitted_ids)):
        raise HTTPException(
            status_code=422,
            detail="Duplicate suggestion IDs",
        )

    if set(submitted_ids) != expected_ids:
        raise HTTPException(
            status_code=422,
            detail="You must review every suggestion",
        )

    decisions = {
        item.suggestion_id: item.decision
        for item in request.decisions
    }

    try:
        updated = await update_optimization_review(
            optimization_id,
            decisions,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=409,
            detail=str(exc),
        ) from exc

    return {
        "optimization_id": updated["_id"],
        "status": updated["status"],
        "suggestions": updated["suggestions"],
        "missing_skills": updated["missing_skills"],
        "summary": updated["summary"],
    }



async def create_approved_resume_version(
    optimization_id: str,
):
    optimization = await get_optimization(optimization_id)

    if optimization is None:
        raise HTTPException(
            status_code=404,
            detail="Optimization not found",
        )

    if optimization["status"] != "reviewed":
        raise HTTPException(
            status_code=409,
            detail="Please review suggestions first",
        )

    # Use the resume snapshot captured during optimization.
    resume_text = optimization["original_resume_text"]

    approved_suggestions = [
        suggestion
        for suggestion in optimization["suggestions"]
        if suggestion["decision"] == "approve"
    ]

    # Check all source strings before modifying the text.
    for suggestion in approved_suggestions:
        original = suggestion["original_text"]

        if not original or resume_text.count(original) != 1:
            raise HTTPException(
                status_code=422,
                detail="Approved suggestion has an ambiguous or missing source",
            )

    # Reject overlapping replacements before applying them.
    ranges = []

    for suggestion in approved_suggestions:
        original = suggestion["original_text"]
        start = resume_text.index(original)

        ranges.append(
            (start, start + len(original), suggestion["suggested_text"])
        )

    ranges.sort(key=lambda item: item[0])

    for previous, current in zip(ranges, ranges[1:]):
        if previous[1] > current[0]:
            raise HTTPException(
                status_code=422,
                detail="Approved suggestions contain overlapping changes",
            )

    # Apply from right to left to preserve source offsets.
    for start, end, replacement in reversed(ranges):
        resume_text = (
            resume_text[:start]
            + replacement
            + resume_text[end:]
        )

    version = await save_resume_version(
        optimization_id=optimization_id,
        resume_text=resume_text,
    )

    return {
        "version_id": version["_id"],
        "optimization_id": version["optimization_id"],
        "resume_text": version["resume_text"],
    }


async def fetch_resume_version(version_id: str):
    version = await get_resume_version(version_id)

    if version is None:
        raise HTTPException(
            status_code=404,
            detail="Resume version not found",
        )

    return {
        "version_id": version["_id"],
        "optimization_id": version["optimization_id"],
        "resume_text": version["resume_text"],
    }